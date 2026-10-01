import os
import re
from docx import Document
from docx.shared import Pt, Cm, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_page_margins(section):
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.left_margin = Cm(3.5)
    section.top_margin = Cm(2.5)
    section.right_margin = Cm(1.25)
    section.bottom_margin = Cm(1.25)

def add_page_number(run):
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = "PAGE"
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')
    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)

def set_page_number_format(section, fmt='lowerRoman'):
    sectPr = section._sectPr
    pgNumType = OxmlElement('w:pgNumType')
    pgNumType.set(qn('w:fmt'), fmt)
    pgNumType.set(qn('w:start'), '1')
    sectPr.append(pgNumType)

def add_toc(paragraph):
    run = paragraph.add_run()
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = 'TOC \\o "1-3" \\h \\z \\u'
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')
    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)

doc = Document()
style = doc.styles['Normal']
font = style.font
font.name = 'Times New Roman'
font.size = Pt(12)
style.paragraph_format.line_spacing = 1.5

def add_paragraph(text, size=12, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, spacing=1.5):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.line_spacing = spacing
    if text:
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = RGBColor(0,0,0)
    return p

def style_heading(h, size):
    for run in h.runs:
        run.font.name = 'Times New Roman'
        run.font.size = Pt(size)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0,0,0)

# SECTION 1: Front Matter
section_1 = doc.sections[0]
set_page_margins(section_1)
set_page_number_format(section_1, 'lowerRoman')

footer_1 = section_1.footer
footer_1.is_linked_to_previous = False
footer_p = footer_1.paragraphs[0]
footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_page_number(footer_p.add_run())

for _ in range(5): doc.add_paragraph()
add_paragraph("SUMMER INTERNSHIP REPORT", size=24, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_paragraph()
add_paragraph("DOCFORGE", size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_paragraph()
doc.add_paragraph()
add_paragraph("Submitted in partial fulfillment of the Requirements for the award of Degree of Bachelor of Technology in Artificial Intelligence and Machine Learning (AI/ML)", size=12, align=WD_ALIGN_PARAGRAPH.CENTER)
for _ in range(3): doc.add_paragraph()
add_paragraph("Submitted By", size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
add_paragraph("Name: Harshil Srivastav\nUniversity Roll No.: 35817711624\nSemester: 5th\nAcademic Year: 2026-27", size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

doc.add_page_break()
for _ in range(2): doc.add_paragraph()
add_paragraph("DECLARATION", size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_paragraph()
add_paragraph('I hereby declare that the Internship Report entitled "DocForge" is an authentic record of my own work as requirements of Internship during the period from 15.06.2026 to 31.07.2026 for the award of degree of B.Tech. (AIML) School of Engineering and Technology, VIPS-TC.', size=12, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
doc.add_paragraph()
doc.add_paragraph()
add_paragraph("(Signature of student)\nHarshil Srivastav\n35817711624\nDate: 25 September 2026", size=12)

doc.add_page_break()
for _ in range(2): doc.add_paragraph()
add_paragraph("CERTIFICATE", size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_paragraph()
add_paragraph("[INSERT C-DAC CERTIFICATE IMAGE HERE]", size=12, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_paragraph()
add_paragraph("This is to certify that Mr. Harshil Srivastav has completed Internship during the period from 15.06.2026 to 31.07.2026 in our Organization / Industry as a Partial Fulfillment of Degree of Bachelor of Technology in AIML. He was trained in the field of Artificial Intelligence and Machine Learning / Document Forgery Detection.", size=12, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
doc.add_paragraph()
doc.add_paragraph()
add_paragraph("Signature & Seal of Training Manager\nMr. Tushar Kanti Patnaik, Scientist F", size=12)

doc.add_page_break()
for _ in range(2): doc.add_paragraph()
add_paragraph("ACKNOWLEDGEMENT", size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_paragraph()
add_paragraph("I would like to express my profound gratitude to the Centre for Development of Advanced Computing (C-DAC), Noida, for providing me with the invaluable opportunity to pursue my summer internship. This experience has significantly enriched my practical understanding of machine learning and computer vision.", size=12, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
add_paragraph("I am deeply indebted to my project guide, Mr. Tushar Kanti Patnaik, Scientist F, for his continuous guidance, encouragement, and technical mentorship throughout the development of the \"DocForge\" project. His insights were instrumental in navigating the complexities of deep learning pipelines.", size=12, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
add_paragraph("I also extend my sincere thanks to the faculty and administration of Vivekananda Institute of Professional Studies - Technical Campus (VIPS-TC), Pitampura, for equipping me with the foundational knowledge required to undertake this training and for their continuous academic support. Finally, I would like to thank my family and peers for their unwavering encouragement throughout this endeavor.", size=12, align=WD_ALIGN_PARAGRAPH.JUSTIFY)

doc.add_page_break()
for _ in range(2): doc.add_paragraph()
add_paragraph("ABOUT COMPANY", size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_paragraph()
add_paragraph("The Centre for Development of Advanced Computing (C-DAC) is the premier research and development organization of the Ministry of Electronics and Information Technology (MeitY), Government of India. Established to advance the country's IT and electronics sector, C-DAC engages in cutting-edge research in areas such as High-Performance Computing (HPC), Artificial Intelligence, Cyber Security, and Software Technologies. The Noida facility, located in the Institutional Area of Sector 62, actively drives technological innovation, education, and specialized training programs designed to bridge the gap between academic learning and industry demands.", size=12, align=WD_ALIGN_PARAGRAPH.JUSTIFY)

doc.add_page_break()
for _ in range(2): doc.add_paragraph()
add_paragraph("TABLE OF CONTENTS", size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_paragraph()
add_toc(doc.add_paragraph())
add_paragraph("(To update this TOC, right click here in Word and select 'Update Field')", size=10, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER)

# SECTION 2: Main Chapters
section_2 = doc.add_section(WD_SECTION.NEW_PAGE)
set_page_margins(section_2)
set_page_number_format(section_2, 'decimal')
footer_2 = section_2.footer
footer_2.is_linked_to_previous = False
footer_p2 = footer_2.paragraphs[0]
footer_p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_page_number(footer_p2.add_run())

CONTENT = [
    ("# CHAPTER 1: INTRODUCTION TO PROJECT", 1, 14),
    ("## 1.1 Problem Statement", 2, 12),
    ("The advent of the digital era has ushered in an unprecedented reliance on digital documents across nearly every facet of modern society. From legal contracts, identity documents, and financial records to academic transcripts and medical histories, the digitization of information has vastly improved the efficiency of storage, retrieval, and transmission.", 0, 12),
    ("However, this convenience is intrinsically accompanied by a significant vulnerability: the ease with which digital documents can be manipulated or forged. Historically, document forgery involved physical alterations—erasing ink, forging signatures by hand, or carefully pasting elements from one document to another. The verification of such documents, while challenging, relied on the physical properties of the paper, ink analysis, and the keen eye of trained forensic document examiners. In contrast, modern digital forgery leverages sophisticated image editing software like Adobe Photoshop, GIMP, and increasingly, generative artificial intelligence models. These tools democratize the ability to alter digital images, allowing even novice users to seamlessly perform copy-move, splicing, and retouching operations without leaving visible traces to the naked eye.", 0, 12),
    ("Copy-move forgery involves copying a specific region of a document and pasting it into another region of the same document, often to conceal compromising information or to artificially duplicate signatures and stamps. Splicing, on the other hand, involves merging elements from two entirely different source documents to create a single composite forgery. Retouching involves subtle pixel-level modifications, such as altering dates or financial figures, without moving large blocks of pixels.", 0, 12),
    ("The limitations of human verification in detecting these sophisticated digital manipulations are profound. Human visual perception is generally ill-equipped to identify subtle statistical anomalies, pixel-level discrepancies, or complex lighting inconsistencies introduced during the forgery process. When a digital document is heavily compressed (e.g., JPEG compression), the resulting compression artifacts can further mask the high-frequency edge anomalies introduced by splicing or copy-move operations, making manual detection nearly impossible.", 0, 12),
    ("Manual inspection is not only highly subjective and error-prone but also extremely time-consuming and unscalable in environments where thousands of documents are processed daily, such as immigration checkpoints, insurance claim processing centers, or large-scale financial institutions. Consequently, there is an urgent and critical need for automated, intelligent systems capable of detecting document manipulations with high accuracy and reliability.", 0, 12),
    
    ("## 1.2 Objectives", 2, 12),
    ("The primary objective of the DocForge project is to design, develop, and evaluate a robust, state-of-the-art deep learning system capable of automatically detecting and localizing digital manipulations in scanned and digital-born documents. Specifically, the objectives are delineated as follows:", 0, 12),
    ("1. To engineer an automated detection pipeline: To transition away from manual, subjective verification processes by developing an automated pipeline that can ingest digital documents and output a pixel-level localization map indicating forged regions. This map will highlight the exact spatial coordinates of the manipulation.", 0, 12),
    ("2. To design a robust feature extraction framework: To leverage advanced Convolutional Neural Networks (CNNs) capable of extracting both high-level semantic features and low-level noise artifacts. High-level features are essential for understanding the context (e.g., text vs background), while low-level features are crucial for identifying the physical discrepancies introduced by forgery operations.", 0, 12),
    ("3. To implement an effective learning strategy: To train the system on diverse and complex datasets comprising both authentic and manipulated documents, ensuring the model generalizes well to unseen forgery techniques. The learning strategy must account for severe class imbalances, as forged pixels typically represent less than 5% of the total document area.", 0, 12),
    ("4. To provide a scalable and efficient solution: To ensure that the resulting system, while highly accurate, maintains computational efficiency suitable for deployment in real-world, high-throughput environments.", 0, 12),
    
    ("## 1.3 Scope", 2, 12),
    ("The scope of the DocForge project is expansive, given the ubiquitous nature of digital documents. The real-world applications of a robust document forgery detection system span multiple critical sectors. In the Legal Sector, the integrity of digital evidence is paramount in legal proceedings. Automated forgery detection can verify the authenticity of submitted digital contracts, emails, and photographic evidence, preventing miscarriages of justice predicated on fabricated documentation.", 0, 12),
    ("Within the Financial Sector, financial institutions are frequent targets for fraud involving manipulated bank statements, identity cards (KYC documents), and checks. An automated system can act as a crucial first line of defense during the onboarding and verification processes, significantly reducing financial fraud and protecting institutional assets.", 0, 12),
    ("In Administrative and Government Sectors, governments process millions of documents ranging from passports and visas to tax records. Ensuring the authenticity of these documents is vital for national security and the proper functioning of administrative apparatuses. Furthermore, in the Insurance Industry, fraudulent insurance claims often involve manipulated images of accidents or damaged property. Automated detection can flag inconsistencies, streamlining the claims process and mitigating immense financial losses due to organized fraud.", 0, 12),
    
    ("# CHAPTER 2: TOOLS & TECHNOLOGY USED", 1, 14),
    ("## 2.1 Deep Learning Fundamentals", 2, 12),
    ("The foundation of the DocForge system lies in Deep Learning, a subfield of Machine Learning inspired by the structure and function of the human brain's neural networks. Deep learning models, specifically deep artificial neural networks, have revolutionized the field of Computer Vision by achieving unprecedented performance in tasks such as image classification, object detection, and semantic segmentation.", 0, 12),
    
    ("### 2.1.1 Artificial Neural Networks (ANNs)", 3, 12),
    ("At the core of deep learning is the Artificial Neural Network. An ANN consists of interconnected nodes, or \"neurons\", organized into distinct layers: an input layer, one or more hidden layers, and an output layer. Each connection between these neurons has an associated \"weight\" and \"bias\". During the forward propagation phase, data is fed into the input layer, multiplied by the weights, summed with the biases, and passed through an activation function.", 0, 12),
    ("Mathematically, the pre-activation output of a neuron 'j' can be described as Z_j = sum(W_ij * X_i) + b_j, where W_ij represents the weight from neuron 'i' to neuron 'j', X_i is the input from neuron 'i', and b_j is the bias term for neuron 'j'. The final output of the neuron is then calculated by applying a non-linear activation function to Z_j. This hierarchical stacking of mathematical transformations allows the network to approximate virtually any continuous function, a property known as the Universal Approximation Theorem.", 0, 12),
    
    ("### 2.1.2 Activation Functions", 3, 12),
    ("Activation functions are fundamental to the operation of deep neural networks. Without them, the entire network—no matter how many layers deep—would mathematically collapse into a single linear regression model, entirely incapable of learning the complex, highly non-linear patterns present in digital imagery and forgery artifacts.", 0, 12),
    ("1. Rectified Linear Unit (ReLU): ReLU is the most prevalent activation function in modern CNNs, defined as f(x) = max(0, x). It allows positive values to pass through unchanged while clamping negative values to zero. This simplicity makes ReLU computationally efficient and helps mitigate the vanishing gradient problem, allowing deep networks to converge much faster than older functions like Sigmoid.", 0, 12),
    ("2. Sigmoid Function: The Sigmoid function, defined as f(x) = 1 / (1 + exp(-x)), squashes input values into a smooth range between 0 and 1. While historically popular, it suffers from severe vanishing gradients at the extremes of its domain, making it unsuitable for hidden layers in deep networks. However, it remains essential in the final output layer of the DocForge architecture to transform the raw model logits into standardized probability maps indicating the likelihood of a pixel being forged.", 0, 12),
    
    ("### 2.1.3 Backpropagation and Gradient Descent", 3, 12),
    ("Training the network to accurately detect forgeries requires an iterative optimization algorithm to find the optimal set of weights and biases. The standard approach involves defining a Loss Function, which quantifies the error between the model's predictions and the actual ground-truth forgery mask. The goal of training is to globally minimize this loss.", 0, 12),
    ("Backpropagation, or backward propagation of errors, is the algorithm used to efficiently compute the gradient (the multi-dimensional derivative) of the loss function with respect to every single weight in the network. It achieves this by applying the chain rule of calculus sequentially from the output layer back to the input layer. These gradients indicate the precise direction and magnitude required to adjust each weight to reduce the overall error.", 0, 12),
    ("Stochastic Gradient Descent (SGD) and its modern variants, such as the Adam (Adaptive Moment Estimation) optimizer, utilize these gradients to iteratively update the network parameters. Adam computes individual adaptive learning rates for different parameters from estimates of first and second moments of the gradients, providing highly stable and rapid convergence even on complex, noisy loss landscapes like those encountered in pixel-level anomaly detection.", 0, 12),

    ("## 2.2 Convolutional Neural Networks (CNNs)", 2, 12),
    ("While traditional dense neural networks are powerful, they are highly inefficient for processing high-resolution image data. A 256x256 RGB image contains nearly 200,000 individual pixel values; a fully connected layer attempting to process this would require billions of parameters, leading to immediate memory exhaustion and severe statistical overfitting. Convolutional Neural Networks (CNNs) solve this problem by leveraging the spatial hierarchy of image data.", 0, 12),
    
    ("### 2.2.1 The Convolution Operation", 3, 12),
    ("The defining operation of a CNN is the convolution. Instead of connecting every input pixel to every neuron, a CNN utilizes small, learnable matrices called filters (or kernels)—typically 3x3 or 5x5 in size. During the forward pass, these filters slide (or convolve) across the width and height of the input image. At each spatial position, a dot product is computed between the filter's weights and the specific patch of the input image it is currently hovering over.", 0, 12),
    ("This sliding window approach produces a 2D activation map (or feature map) that highlights the spatial locations where the filter detected its specific visual pattern. Early layers in the network learn to detect primitive features like horizontal edges, vertical edges, and color gradients. Deeper layers combine these primitive features to detect complex textures, shapes, and semantic objects. Because the same filter is applied across the entire image (Weight Sharing), the network requires vastly fewer parameters and inherently achieves translation invariance—if a forged artifact moves to a different location in the document, the same filter will still detect it.", 0, 12),
    
    ("### 2.2.2 Padding and Stride", 3, 12),
    ("Two critical hyperparameters control the spatial dimensions of the output feature maps produced by convolutional layers: padding and stride.", 0, 12),
    ("Stride dictates the step size the filter takes as it slides across the image. A stride of 1 means the filter moves one pixel at a time, resulting in a dense feature map. A stride of 2 means the filter skips every other pixel, effectively halving the spatial dimensions of the output and accelerating computation.", 0, 12),
    ("Padding involves adding artificial pixels (typically zeros, known as Zero-Padding) around the perimeter of the input image before applying the convolution. Without padding, a convolution operation naturally shrinks the spatial dimensions of the image, and pixels located at the edges are under-represented because the filter passes over them fewer times than central pixels. By applying 'Same' padding, the network preserves the original spatial resolution throughout the convolutional layers, which is crucial for dense prediction tasks like forgery localization.", 0, 12),
    
    ("### 2.2.3 Pooling and Upsampling", 3, 12),
    ("Pooling layers are periodically inserted between convolutional layers to aggressively downsample the spatial dimensions of the feature maps. Max Pooling is the most common technique; it slides a window (e.g., 2x2) over the feature map and simply outputs the maximum value within that window. This radically reduces the computational burden and forces the network to learn spatially invariant features.", 0, 12),
    ("However, for forgery localization, the network must eventually output a mask that matches the exact original resolution of the input document. To achieve this, the architecture employs an encoder-decoder structure. The encoder uses convolutions and pooling to compress the image and extract deep semantic features. The decoder uses Upsampling operations (such as Bilinear Interpolation or Transposed Convolutions) to progressively rebuild the spatial resolution, fusing these deep semantic features with high-resolution structural features skipped over from the encoder.", 0, 12),
    
    ("## 2.3 Attention Mechanisms", 2, 12),
    ("A critical enhancement to modern CNN architectures is the integration of Attention Mechanisms. In standard CNNs, all channels and all spatial locations within a feature map are treated with equal importance. However, in forgery detection, the vast majority of the document consists of authentic background paper texture and text, which acts as noise relative to the actual task. Attention mechanisms allow the network to dynamically weigh the importance of features, selectively suppressing background noise and amplifying anomalous forgery artifacts.", 0, 12),
    
    ("### 2.3.1 Spatial Attention", 3, 12),
    ("Spatial Attention focuses on determining *where* the relevant information is located. It operates by aggregating feature information across the channel dimension using both Max Pooling and Average Pooling. These pooled representations are concatenated and passed through a standard convolutional layer followed by a Sigmoid activation. This generates a 2D spatial weight map. When this map is multiplied element-wise with the original feature map, it acts as a filter, explicitly telling the network to ignore pristine background regions and focus intense computational resources on suspected manipulation boundaries.", 0, 12),
    
    ("### 2.3.2 Channel Attention", 3, 12),
    ("Channel Attention focuses on determining *what* features are most important. A deep feature map might consist of 256 or 512 separate channels, each representing a different learned filter (e.g., one channel might track edge noise, while another tracks ink color). Channel attention aggregates the spatial information of each channel using global pooling operations and processes it through a small Multi-Layer Perceptron (MLP). The MLP outputs a 1D vector of weights that scales each individual channel. This mechanism allows the network to dynamically prioritize filters that detect high-frequency compression artifacts or noise inconsistencies, which are the primary hallmarks of digital splicing.", 0, 12),

    ("## 2.4 Software Frameworks", 2, 12),
    ("### 2.4.1 Python and PyTorch", 3, 12),
    ("The entire DocForge pipeline is implemented in Python, leveraging the PyTorch deep learning framework. PyTorch fundamentally operates on Tensors, which are multi-dimensional arrays optimized for GPU acceleration via NVIDIA's CUDA platform.", 0, 12),
    ("A defining feature of PyTorch utilized extensively in this project is its Autograd engine. During the forward pass, Autograd dynamically constructs a Directed Acyclic Graph (DAG) representing the sequence of mathematical operations applied to the tensors. When the loss is calculated, calling the `.backward()` method triggers the Autograd engine to traverse this DAG in reverse, automatically calculating and accumulating the exact analytical gradients for every single learnable parameter in the architecture. This dynamic 'define-by-run' execution model is highly advantageous for complex architectures featuring custom attention modules and dynamic loss scaling.", 0, 12),
    
    ("### 2.4.2 OpenCV and Data Processing", 3, 12),
    ("OpenCV (Open Source Computer Vision Library) forms the backbone of the system's data ingestion pipeline. Before PyTorch can analyze a document, OpenCV handles the low-level memory operations to read compressed JPEGs and PNGs from disk into NumPy matrices.", 0, 12),
    ("OpenCV is heavily utilized to process the ground-truth binary masks provided in the academic datasets. For structural analysis, OpenCV's Canny Edge Detector is applied to the binary forgery masks. The Canny algorithm calculates image gradients to identify regions of rapid intensity change, effectively tracing the exact perimeter of the manipulated region. This allows the neural network to explicitly calculate loss metrics based not just on the area of the forgery, but on the precise accuracy of the predicted forgery boundary.", 0, 12),
    
    ("# CHAPTER 3: SYSTEM ARCHITECTURE & IMPLEMENTATION", 1, 14),
    ("## 3.1 Data Pipeline and Augmentation", 2, 12),
    ("A deep learning model is only as robust as the data it is trained on. To prevent the high-capacity CNN from simply memorizing the training dataset (overfitting), a highly aggressive data augmentation pipeline was implemented.", 0, 12),
    ("During training, the data loader dynamically applies geometric transformations on the fly. These include random horizontal and vertical flips, and random continuous rotations between -15 and 15 degrees. When rotating, the system must carefully interpolate the image pixels while strictly using nearest-neighbor interpolation on the binary mask to prevent anti-aliasing artifacts from creating false 'partial' forgery labels.", 0, 12),
    ("Photometric distortions are equally critical. The pipeline applies random Gaussian blurring, brightness shifting, and contrast scaling. This simulates real-world conditions where a document might be scanned by a low-quality flatbed scanner, photographed by a smartphone under poor lighting, or heavily compressed for email transmission. By forcing the network to detect forgeries under these adverse simulated conditions, the final model achieves significantly higher generalization accuracy in real-world deployment.", 0, 12),
    
    ("## 3.2 Abstracted Architecture Description", 2, 12),
    ("Due to the proprietary nature of the research, the exact novel architectural formulations are abstracted. The system operates on a dual-branch philosophy designed to capture both semantic layout and low-level noise.", 0, 12),
    ("The primary Feature Extraction Backbone processes the RGB document image to understand the semantic context. Simultaneously, a secondary analytical pathway evaluates high-frequency noise residuals. In authentic documents, camera sensor noise and JPEG compression artifacts are generally uniform across the entire image. However, when a region is spliced from a different image, or when a region is smoothed using a retouching tool, this uniform noise distribution is violently disrupted.", 0, 12),
    ("The Custom Attention-Based Fusion Module is responsible for reconciling these two streams of data. By cross-referencing semantic boundaries with noise inconsistencies, the network can differentiate between a legitimate physical edge (e.g., the edge of a photograph printed on an ID card) and a malicious digital edge (a forged photograph pasted over the original). The fused representations are decoded and passed through a final 1x1 convolution with a Sigmoid activation, outputting a precise heat-map of suspected manipulation.", 0, 12),
    
