import os
import re
from docx import Document
from docx.shared import Pt, Cm, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_page_margins(section):
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.left_margin = Cm(3.5)
    section.top_margin = Cm(2.5)
    section.right_margin = Cm(1.25)
    section.bottom_margin = Cm(1.25)

def add_page_number(run):
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = "PAGE"
    
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')
    
    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')
    
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)

def set_page_number_format(section, fmt='lowerRoman'):
    sectPr = section._sectPr
    pgNumType = OxmlElement('w:pgNumType')
    pgNumType.set(qn('w:fmt'), fmt)
    # Start at 1
    pgNumType.set(qn('w:start'), '1')
    sectPr.append(pgNumType)

doc = Document()

# Set up global style
style = doc.styles['Normal']
font = style.font
font.name = 'Times New Roman'
font.size = Pt(12)
style.paragraph_format.line_spacing = 1.5

def add_paragraph(text, size=12, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, spacing=1.5):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.line_spacing = spacing
    
    if text:
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(size)
        run.font.bold = bold
    return p

# Section 1: Front Matter
section_1 = doc.sections[0]
set_page_margins(section_1)
set_page_number_format(section_1, 'lowerRoman')

# Footer for Section 1
footer_1 = section_1.footer
footer_1.is_linked_to_previous = False
footer_p = footer_1.paragraphs[0]
footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer_run = footer_p.add_run()
add_page_number(footer_run)

# Cover Page
for _ in range(5): doc.add_paragraph()
add_paragraph("SUMMER INTERNSHIP REPORT", size=24, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_paragraph()
add_paragraph("DOCFORGE", size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_paragraph()
doc.add_paragraph()
add_paragraph("Submitted in partial fulfillment of the Requirements for the award of Degree of Bachelor of Technology in Artificial Intelligence and Machine Learning (AI/ML)", size=12, align=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_paragraph()
doc.add_paragraph()
doc.add_paragraph()
add_paragraph("Submitted By", size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
add_paragraph("Name: Harshil Srivastav\nUniversity Roll No.: 35817711624\nSemester: 5th\nAcademic Year: 2026-27", size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

doc.add_page_break()

# Declaration
for _ in range(2): doc.add_paragraph()
add_paragraph("DECLARATION", size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_paragraph()
add_paragraph('I hereby declare that the Internship Report entitled "DocForge" is an authentic record of my own work as requirements of Internship during the period from 15.06.2026 to 31.07.2026 for the award of degree of B.Tech. (AIML) School of Engineering and Technology, VIPS-TC.', size=12, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
doc.add_paragraph()
doc.add_paragraph()
add_paragraph("(Signature of student)\nHarshil Srivastav\n35817711624\nDate: 25 September 2026", size=12)

doc.add_page_break()

# Certificate
for _ in range(2): doc.add_paragraph()
add_paragraph("CERTIFICATE", size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_paragraph()
add_paragraph("[INSERT C-DAC CERTIFICATE IMAGE HERE]", size=12, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_paragraph()
add_paragraph("This is to certify that Mr. Harshil Srivastav has completed Internship during the period from 15.06.2026 to 31.07.2026 in our Organization / Industry as a Partial Fulfillment of Degree of Bachelor of Technology in AIML. He was trained in the field of Artificial Intelligence and Machine Learning / Document Forgery Detection.", size=12, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
doc.add_paragraph()
doc.add_paragraph()
add_paragraph("Signature & Seal of Training Manager\nMr. Tushar Kanti Patnaik, Scientist F", size=12)

doc.add_page_break()

# Acknowledgement
for _ in range(2): doc.add_paragraph()
add_paragraph("ACKNOWLEDGEMENT", size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_paragraph()
add_paragraph("I would like to express my profound gratitude to the Centre for Development of Advanced Computing (C-DAC), Noida, for providing me with the invaluable opportunity to pursue my summer internship. This experience has significantly enriched my practical understanding of machine learning and computer vision.", size=12, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
add_paragraph("I am deeply indebted to my project guide, Mr. Tushar Kanti Patnaik, Scientist F, for his continuous guidance, encouragement, and technical mentorship throughout the development of the \"DocForge\" project. His insights were instrumental in navigating the complexities of deep learning pipelines.", size=12, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
add_paragraph("I also extend my sincere thanks to the faculty and administration of Vivekananda Institute of Professional Studies - Technical Campus (VIPS-TC), Pitampura, for equipping me with the foundational knowledge required to undertake this training and for their continuous academic support. Finally, I would like to thank my family and peers for their unwavering encouragement throughout this endeavor.", size=12, align=WD_ALIGN_PARAGRAPH.JUSTIFY)

doc.add_page_break()

# About Company
for _ in range(2): doc.add_paragraph()
add_paragraph("ABOUT COMPANY", size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_paragraph()
add_paragraph("The Centre for Development of Advanced Computing (C-DAC) is the premier research and development organization of the Ministry of Electronics and Information Technology (MeitY), Government of India. Established to advance the country's IT and electronics sector, C-DAC engages in cutting-edge research in areas such as High-Performance Computing (HPC), Artificial Intelligence, Cyber Security, and Software Technologies. The Noida facility, located in the Institutional Area of Sector 62, actively drives technological innovation, education, and specialized training programs designed to bridge the gap between academic learning and industry demands.", size=12, align=WD_ALIGN_PARAGRAPH.JUSTIFY)

doc.add_page_break()

# TOC place holders
for _ in range(2): doc.add_paragraph()
add_paragraph("TABLE OF CONTENTS", size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
add_paragraph("[INSERT AUTO-GENERATED TABLE OF CONTENTS HERE]", size=12, align=WD_ALIGN_PARAGRAPH.CENTER)
